# -*- coding: utf-8 -*-
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import openpyxl


def determine_row_result(row_summary: dict[str, Any]) -> tuple[str, str]:
    """
    根据 row_summary 判定行级最终结果和失败原因。

    返回 (result, reason):
      result: yes / suspected / no / failed
      reason: 失败阶段描述，无失败时为空字符串
    """
    status = row_summary.get("status", "")
    task_summaries = row_summary.get("task_summaries", [])
    analysis = row_summary.get("analysis", {})

    # 优先从 row_summary.analysis.counts 获取聚合结果
    row_counts = analysis.get("counts", {})
    yes_count = int(row_counts.get("yes", 0))
    suspected_count = int(row_counts.get("suspected", 0))
    no_count = int(row_counts.get("no", 0))

    # 从 task_summaries 收集失败阶段信息
    failed_stages: list[str] = []
    for task in task_summaries:
        task_status = task.get("status", "")
        ai_status = task.get("ai_status")
        if task_status in ("download_failed",):
            failed_stages.append("下载失败")
        elif task_status in ("decode_failed", "worker_failed", "missing_target_ts"):
            failed_stages.append("解码失败")
        elif ai_status in ("failed", "partial_failed"):
            failed_stages.append("AI推理失败")

    has_valid_results = (yes_count + suspected_count + no_count) > 0

    # 判定结果
    if not has_valid_results:
        reason = "; ".join(sorted(set(failed_stages))) if failed_stages else "全部任务失败"
        return "failed", reason

    if yes_count > 0:
        result = "yes"
    elif suspected_count > 0:
        result = "suspected"
    else:
        result = "no"

    reason = "; ".join(sorted(set(failed_stages))) if failed_stages else ""
    return result, reason


def run_writeback(config_path: str | Path) -> dict[str, Any]:
    """
    读取输出目录中的 row_summary，回填结果到 Excel 最后两列。
    """
    config_path = Path(config_path)
    with open(config_path, "r", encoding="utf-8") as f:
        cfg = json.load(f)

    excel_cfg = cfg.get("excel", {})
    excel_path = Path(excel_cfg["path"])
    sheet_name = excel_cfg.get("sheet", 0)
    header_row = int(excel_cfg.get("header", 0)) + 1  # openpyxl 1-indexed

    output_root = Path(
        cfg.get("output", {}).get("root_dir")
        or excel_cfg.get("output_dir")
        or "outputs"
    )

    if not excel_path.exists():
        raise FileNotFoundError(f"Excel 文件不存在: {excel_path}")

    wb = openpyxl.load_workbook(str(excel_path))
    if isinstance(sheet_name, int):
        ws = wb.worksheets[sheet_name]
    else:
        ws = wb[sheet_name]

    # 定位最后两列（在已有数据列之后）
    max_col = ws.max_column
    result_col_name = "workflow_result"
    reason_col_name = "workflow_reason"

    # 检查是否已有回填列（支持重复执行覆盖）
    result_col = None
    reason_col = None
    for col_idx in range(1, max_col + 1):
        cell_value = ws.cell(row=header_row, column=col_idx).value
        if cell_value == result_col_name:
            result_col = col_idx
        elif cell_value == reason_col_name:
            reason_col = col_idx

    if result_col is None:
        result_col = max_col + 1
        ws.cell(row=header_row, column=result_col, value=result_col_name)
    if reason_col is None:
        reason_col = result_col + 1
        ws.cell(row=header_row, column=reason_col, value=reason_col_name)

    # 收集所有 row_summary
    results: list[dict[str, Any]] = []
    if not output_root.exists():
        print(f"[WARN] 输出目录不存在: {output_root}")
        return {"written": 0, "results": results}

    for row_dir in sorted(output_root.iterdir()):
        if not row_dir.is_dir():
            continue
        summary_path = row_dir / "row_summary.json"
        meta_path = row_dir / "row_meta.json"

        if not summary_path.exists() and not meta_path.exists():
            continue

        # 获取 excel_row
        excel_row = None
        if meta_path.exists():
            with open(meta_path, "r", encoding="utf-8") as f:
                meta = json.load(f)
            excel_row = meta.get("excel_row")

        if excel_row is None:
            print(f"[WARN] {row_dir.name}: 无法确定 excel_row，跳过回填")
            continue

        if summary_path.exists():
            with open(summary_path, "r", encoding="utf-8") as f:
                row_summary = json.load(f)
        else:
            row_summary = {"status": "no_summary", "tasks": []}

        result, reason = determine_row_result(row_summary)

        ws.cell(row=int(excel_row), column=result_col, value=result)
        ws.cell(row=int(excel_row), column=reason_col, value=reason)

        entry = {
            "row_dir": row_dir.name,
            "excel_row": excel_row,
            "result": result,
            "reason": reason,
        }
        results.append(entry)
        print(f"[WRITEBACK] row={excel_row} result={result} reason={reason or '-'}")

    wb.save(str(excel_path))
    print(f"[INFO] 回填完成: {len(results)} 行已写入 {excel_path}")

    # 保存回填汇总
    summary = {"excel_path": str(excel_path), "written": len(results), "results": results}
    summary_path = output_root / "writeback_summary.json"
    with open(summary_path, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    return summary


if __name__ == "__main__":
    import sys
    if len(sys.argv) != 2:
        print("用法: python writeback.py config.json")
        sys.exit(1)
    run_writeback(sys.argv[1])
